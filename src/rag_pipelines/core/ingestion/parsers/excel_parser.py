import openpyxl
from io import BytesIO
from rag_pipelines.core.models import ParsedChunk
from rag_pipelines.core.ingestion.parsers.base import BaseParser


class ExcelParser(BaseParser):
    model_config = {"arbitrary_types_allowed": True}
    def parse(self, file_stream: BytesIO, filename: str) -> list[ParsedChunk]:
        chunks = []
        wb = openpyxl.load_workbook(file_stream, data_only=True)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))

            if not rows:
                continue

            headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]

            for row in rows[1:]:
                if all(v is None for v in row):
                    continue
                content = {headers[i]: row[i] for i in range(len(headers))}
                chunks.append(ParsedChunk(
                    content=content,
                    source=filename,
                    sheet=sheet_name,
                ))

        return chunks
