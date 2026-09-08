from io import BytesIO
from pptx import Presentation
from rag_pipelines.core.models import ParsedChunk
from rag_pipelines.core.ingestion.parsers.base import BaseParser


class PowerpointParser(BaseParser):
    model_config = {"arbitrary_types_allowed": True}
    def parse(self, file_stream: BytesIO, filename: str) -> list[ParsedChunk]:
        chunks = []
        prs = Presentation(file_stream)

        for i, slide in enumerate(prs.slides, start=1):
            title = None
            body_parts: list[str] = []

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if shape.shape_type == 13 or (hasattr(shape, "placeholder_format") and shape.placeholder_format and shape.placeholder_format.idx == 0):
                    title = text
                else:
                    body_parts.append(text)

            if title or body_parts:
                chunks.append(ParsedChunk(
                    content={"title": title, "body": " ".join(body_parts)},
                    source=filename,
                    slide=i,
                ))

        return chunks
